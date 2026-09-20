"""
Convert manuscript/manuscript_rendered.md into a submission-ready .docx for
BMC Medical Informatics and Decision Making.

Handles:
  - Stripping HTML comment blocks and the drafting-only blockquote note
    (SUBMISSION_CHECKLIST.md item 22: "Delete the ... drafting-note comments")
  - Converting [^label] footnote citations to numbered [N] inline markers,
    in order of first appearance, and the "[^label]: text" definition lines
    into a numbered reference list
  - Embedding the three figures from outputs/analysis/ at their caption lines
  - Headings, tables, bold/italic paragraphs, lists (same conventions as
    CMR_QuickSubmit/generate_docx.py)

Run: python3 scripts/generate_submission_docx.py
"""
import re
import os

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MD_PATH = os.path.join(REPO_ROOT, "manuscript", "manuscript_rendered.md")
DOCX_PATH = os.path.join(
    REPO_ROOT, "manuscript", "Senapati_Manuscript_BMC_Medical_Informatics_and_Decision_Making.docx"
)
FIGURE_DIR = os.path.join(REPO_ROOT, "outputs", "analysis")
FIGURE_MAP = {
    "**Figure 1.**": "figure1_frontier.png",
    "**Figure 2.**": "figure2_autonomy_states.png",
    "**Figure 3.**": "figure3_gateway_schematic.png",
}

BODY_FONT = "Times New Roman"
HEADING_FONT = "Times New Roman"
BODY_SIZE = 12
LINE_SPACING = 1.5

FOOTNOTE_REF_RE = re.compile(r"\[\^([a-zA-Z0-9]+)\]")
FOOTNOTE_DEF_RE = re.compile(r"^\[\^([a-zA-Z0-9]+)\]:\s*(.*)")


def set_paragraph_spacing(paragraph, space_before=0, space_after=6, line_spacing=LINE_SPACING):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = Pt(BODY_SIZE * line_spacing)


def add_run_with_inline_formatting(paragraph, text, default_bold=False, default_italic=False,
                                    font_name=BODY_FONT, font_size=BODY_SIZE):
    """Parse **bold**, *italic*, `code`, and [N] superscript citation markers."""
    pattern = re.compile(r"(\*\*\*(.+?)\*\*\*|\*\*(.+?)\*\*|\*(.+?)\*|`(.+?)`|\[(\d+(?:,\s*\d+)*)\])")
    last = 0
    for m in pattern.finditer(text):
        if m.start() > last:
            r = paragraph.add_run(text[last:m.start()])
            r.font.name = font_name
            r.font.size = Pt(font_size)
            r.bold = default_bold
            r.italic = default_italic
        if m.group(2):
            r = paragraph.add_run(m.group(2))
            r.font.name = font_name
            r.font.size = Pt(font_size)
            r.bold = True
            r.italic = True
        elif m.group(3):
            r = paragraph.add_run(m.group(3))
            r.font.name = font_name
            r.font.size = Pt(font_size)
            r.bold = True
        elif m.group(4):
            r = paragraph.add_run(m.group(4))
            r.font.name = font_name
            r.font.size = Pt(font_size)
            r.italic = True
        elif m.group(5):
            r = paragraph.add_run(m.group(5))
            r.font.name = "Courier New"
            r.font.size = Pt(font_size)
        elif m.group(6):
            r = paragraph.add_run(m.group(6))
            r.font.name = font_name
            r.font.size = Pt(font_size)
            r.font.superscript = True
        last = m.end()
    if last < len(text):
        r = paragraph.add_run(text[last:])
        r.font.name = font_name
        r.font.size = Pt(font_size)
        r.bold = default_bold
        r.italic = default_italic


def add_table_from_lines(doc, table_lines):
    rows = []
    for line in table_lines:
        if re.match(r"^\|[-| :]+\|$", line.strip()):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        rows.append(cells)
    if not rows:
        return
    cols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for r_idx, row_data in enumerate(rows):
        row = table.rows[r_idx]
        for c_idx, cell_text in enumerate(row_data):
            if c_idx >= cols:
                continue
            cell = row.cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            add_run_with_inline_formatting(p, cell_text, default_bold=(r_idx == 0), font_size=10)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
    doc.add_paragraph()


def add_figure(doc, caption, img_path):
    if not os.path.exists(img_path):
        print(f"  WARNING: image not found: {img_path}")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_before = Pt(10)
    pf.space_after = Pt(4)
    pf.line_spacing = None
    run = p.add_run()
    run.add_picture(img_path, width=Inches(5.8))

    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run_with_inline_formatting(cap, caption, font_size=10)
    pf2 = cap.paragraph_format
    pf2.space_before = Pt(0)
    pf2.space_after = Pt(14)
    pf2.line_spacing = None


def strip_drafting_notes(text: str) -> str:
    """Remove HTML comment blocks and the rendering-note blockquote."""
    text = re.sub(r"<!--.*?-->\n?", "", text, flags=re.DOTALL)
    lines = text.split("\n")
    out = []
    skipping_blockquote = False
    for line in lines:
        if line.strip().startswith("> **RENDERING NOTE"):
            skipping_blockquote = True
            continue
        if skipping_blockquote:
            if line.strip().startswith(">"):
                continue
            skipping_blockquote = False
        out.append(line)
    return "\n".join(out)


def renumber_footnotes(text: str) -> str:
    """Replace [^label] citation markers with [N] and definitions with N. text.

    Definition lines ("[^label]: text") must be matched against the
    *original* text before the inline-citation substitution runs, since that
    substitution would otherwise also rewrite the "[^label]" at the start of
    a definition line into "[N]", leaving a stray ": text" that no longer
    matches the definition pattern.
    """
    order = []
    for m in FOOTNOTE_REF_RE.finditer(text):
        label = m.group(1)
        if label not in order:
            order.append(label)
    numbers = {label: i + 1 for i, label in enumerate(order)}

    def _ref_sub(m):
        return f"[{numbers[m.group(1)]}]"

    out_lines = []
    for line in text.split("\n"):
        m = FOOTNOTE_DEF_RE.match(line.strip())
        if m:
            label, rest = m.group(1), m.group(2)
            n = numbers.get(label)
            if n:
                out_lines.append(f"{n}. {FOOTNOTE_REF_RE.sub(_ref_sub, rest)}")
                continue
        out_lines.append(FOOTNOTE_REF_RE.sub(_ref_sub, line))
    return "\n".join(out_lines)


_HARD_BREAK_RE = re.compile(r"^(#{1,3}\s|>\s|\|.*\||\*\*.*\*\*$|`.*`$|---$)")
_LIST_START_RE = re.compile(r"^(-\s|\d+\.\s)")


def _is_hard_break(stripped: str) -> bool:
    if not stripped:
        return True
    for marker in FIGURE_MAP:
        if stripped.startswith(marker):
            return True
    if stripped.startswith("Correspondence:") or stripped.startswith("ORCID:"):
        return True
    return bool(_HARD_BREAK_RE.match(stripped))


def _join_wrapped_paragraphs(lines):
    """Merge consecutive plain-text lines into single logical paragraphs,
    including continuation lines of a numbered/bulleted list item.

    The source markdown hard-wraps prose at ~80 columns; without this, each
    wrapped line becomes its own short Word paragraph instead of one flowing
    paragraph or list item.
    """
    out = []
    buf = None
    for raw in lines:
        stripped = raw.strip()
        if _is_hard_break(stripped):
            if buf is not None:
                out.append(buf)
                buf = None
            out.append(raw)
            continue
        if _LIST_START_RE.match(stripped):
            if buf is not None:
                out.append(buf)
            buf = stripped
            continue
        if buf is None:
            buf = stripped
        else:
            buf = buf + " " + stripped
    if buf is not None:
        out.append(buf)
    return out


def build_docx(md_path, docx_path):
    with open(md_path, "r", encoding="utf-8") as f:
        text = f.read()

    text = strip_drafting_notes(text)
    text = renumber_footnotes(text)
    lines = text.split("\n")

    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.25)

    lines = _join_wrapped_paragraphs(lines)

    i = 0
    table_buffer = []

    while i < len(lines):
        raw = lines[i].rstrip("\n")
        stripped = raw.strip()

        if stripped == "---":
            p = doc.add_paragraph()
            set_paragraph_spacing(p, 2, 2)
            i += 1
            continue

        if stripped.startswith("|") and stripped.endswith("|"):
            table_buffer.append(stripped)
            i += 1
            continue
        else:
            if table_buffer:
                add_table_from_lines(doc, table_buffer)
                table_buffer = []

        matched_figure = None
        for marker, fname in FIGURE_MAP.items():
            if stripped.startswith(marker):
                matched_figure = fname
                break
        if matched_figure:
            add_figure(doc, stripped, os.path.join(FIGURE_DIR, matched_figure))
            i += 1
            continue

        if stripped.startswith("# ") and not stripped.startswith("## "):
            text_ = stripped[2:].strip()
            p = doc.add_paragraph()
            r = p.add_run(text_)
            r.font.name = HEADING_FONT
            r.font.size = Pt(16)
            r.font.bold = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_spacing(p, 0, 12)
            i += 1
            continue

        if stripped.startswith("## ") and not stripped.startswith("### "):
            text_ = stripped[3:].strip()
            p = doc.add_paragraph()
            r = p.add_run(text_)
            r.font.name = HEADING_FONT
            r.font.size = Pt(13)
            r.font.bold = True
            set_paragraph_spacing(p, 12, 4)
            i += 1
            continue

        if stripped.startswith("### "):
            text_ = stripped[4:].strip()
            p = doc.add_paragraph()
            r = p.add_run(text_)
            r.font.name = HEADING_FONT
            r.font.size = Pt(12)
            r.font.bold = True
            r.font.italic = True
            set_paragraph_spacing(p, 8, 2)
            i += 1
            continue

        if stripped.startswith("> "):
            text_ = stripped[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.right_indent = Inches(0.5)
            add_run_with_inline_formatting(p, text_, font_size=11)
            for run in p.runs:
                run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            set_paragraph_spacing(p, 2, 2)
            i += 1
            continue

        num_match = re.match(r"^(\d+)\.\s+(.*)", stripped)
        if num_match:
            text_ = num_match.group(2).strip()
            p = doc.add_paragraph(style="List Number")
            add_run_with_inline_formatting(p, text_, font_size=10)
            set_paragraph_spacing(p, 0, 3)
            i += 1
            continue

        if stripped.startswith("- "):
            text_ = stripped[2:].strip()
            p = doc.add_paragraph(style="List Bullet")
            add_run_with_inline_formatting(p, text_)
            set_paragraph_spacing(p, 0, 3)
            i += 1
            continue

        if stripped.startswith("**") and stripped.endswith("**") and not stripped.startswith("***"):
            text_ = stripped.strip("*")
            p = doc.add_paragraph()
            r = p.add_run(text_)
            r.font.name = BODY_FONT
            r.font.size = Pt(BODY_SIZE)
            r.font.bold = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_spacing(p, 0, 2)
            i += 1
            continue

        if stripped.startswith("`") and stripped.endswith("`"):
            text_ = stripped.strip("`")
            p = doc.add_paragraph()
            r = p.add_run(text_)
            r.font.name = "Courier New"
            r.font.size = Pt(BODY_SIZE)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_spacing(p, 0, 2)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        p = doc.add_paragraph()
        add_run_with_inline_formatting(p, stripped)
        set_paragraph_spacing(p, 0, 6)
        i += 1

    if table_buffer:
        add_table_from_lines(doc, table_buffer)

    doc.save(docx_path)
    print(f"Saved: {docx_path}")


if __name__ == "__main__":
    build_docx(MD_PATH, DOCX_PATH)
