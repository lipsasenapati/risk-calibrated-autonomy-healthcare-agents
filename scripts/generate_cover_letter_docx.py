"""
Convert manuscript/cover_letter.md into a .docx for BMC Medical Informatics
and Decision Making submission.

Run: python3 scripts/generate_cover_letter_docx.py
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from generate_submission_docx import (  # noqa: E402
    BODY_FONT,
    BODY_SIZE,
    add_run_with_inline_formatting,
    set_paragraph_spacing,
)
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MD_PATH = os.path.join(REPO_ROOT, "manuscript", "cover_letter.md")
DOCX_PATH = os.path.join(
    REPO_ROOT, "manuscript", "Senapati_Cover_Letter_BMC_Medical_Informatics_and_Decision_Making.docx"
)


def build_docx(md_path, docx_path):
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.25)

    skip_until_hr = False
    buf = None
    started_body = False

    def flush():
        nonlocal buf
        if buf:
            p = doc.add_paragraph()
            add_run_with_inline_formatting(p, buf)
            set_paragraph_spacing(p, 0, 10)
        buf = None

    for raw in lines:
        stripped = raw.strip()

        # Skip the leading title and the italic instruction blockquote, up to
        # the first "---" divider.
        if not started_body:
            if stripped == "---":
                started_body = True
            continue

        if stripped == "---":
            flush()
            continue

        if not stripped:
            flush()
            continue

        if stripped.startswith("**") and stripped.endswith("**"):
            flush()
            p = doc.add_paragraph()
            r = p.add_run(stripped.strip("*"))
            r.font.name = BODY_FONT
            r.font.size = Pt(BODY_SIZE)
            r.bold = True
            set_paragraph_spacing(p, 4, 10)
            continue

        if buf is None:
            buf = stripped
        else:
            buf = buf + " " + stripped

    flush()
    doc.save(docx_path)
    print(f"Saved: {docx_path}")


if __name__ == "__main__":
    build_docx(MD_PATH, DOCX_PATH)
